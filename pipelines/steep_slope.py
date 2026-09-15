"""国土数値情報 A47 急傾斜地崩壊危険区域データのダウンロード。

国土交通省「国土数値情報ダウンロードサイト」から急傾斜地崩壊危険区域データを
都道府県単位でダウンロードし、Parquet に変換して配置する。急傾斜地の崩壊による
災害の防止に関する法律（急傾斜地法）第 3 条に基づき都道府県知事が指定した区域で、
土砂災害防止法の警戒区域（A33）とは根拠法も規制の内容も違う。

利用条件による絞り込み
----------------------
急傾斜地崩壊危険区域データは原典資料を作成した都道府県ごとに公開条件が異なる。
提供元が配布する「各都道府県のデータの利用条件について」（都道府県別の一覧）には
使用許諾が

  オープンデータにて公開 / 条件付き公開 / 条件付き公開（独自の公開条件等） /
  非公開 / 未作成

のいずれかで示されている。ここでは「オープンデータにて公開」の都道府県だけを
取り込む（条件付き公開は成果物への特記事項の明記など県ごとの条件の承諾が前提に
なる）。判定は一覧を毎回取得して行い、列の並びが変わった場合・zip の都道府県が
一覧に無い場合はエラーで停止する。

配布ファイルの形式
------------------
zip には都道府県 1 ファイルの GeoJSON が入る。シェープファイルも同梱されるが、
.dbf が CP932 で DuckDB の ST_Read が読めないため GeoJSON を使う。

ダウンロード一覧には 2020 年度整備分と 2021 年度整備分が並ぶ。都道府県単位の
2021 年度分は同じ都道府県の 2020 年度分を作り直したもので、2020 年度分の
ポリゴン 16,123 のうち 2021 年度分のどのポリゴンとも交わらないものは 11
（青森 4・岐阜 7）しかない。2021 年度分だけを取り、2020 年度分は取らない。
一覧には地方単位のまとめ（都道府県コードの範囲外の 52〜59）も並ぶが、これは
都道府県の zip をまとめた zip なので都道府県コードの一致で落ちる。

データソース: 急傾斜地崩壊危険区域（A47、2021年度）
https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-A47-2021.html
"""

import logging
import os
import re
import shutil
import zipfile
from pathlib import Path
from urllib.parse import urljoin
from urllib.request import Request, urlopen

import duckdb
import openpyxl

from pipelines.download import download

logger = logging.getLogger("pipelines")

PAGE_URL = "https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-A47-2021.html"

# 取り込む整備年度（ファイル名の A47-<yy>）
FILE_PREFIX = "A47-21"

# 都道府県別の公開条件（提供元が配布する Excel）
TERMS_URL = "https://nlftp.mlit.go.jp/ksj/gml/codelist/R6_Terms_of_Use_Steep_slope.xlsx"

# 一覧の列位置（0 始まり）と、その列に入っているはずの見出し
TERMS_COLUMNS = {
    "prefecture_code": (0, "都道府県コード"),
    "prefecture_name": (1, "都道府県"),
    "disclosure": (2, "使用許諾"),
    "condition": (3, "条件"),
    "condition_detail": (4, "条件詳細"),
}

# 取り込む使用許諾。これ以外（条件付き公開・非公開・未作成）は除外する
OPEN_DISCLOSURE = "オープンデータにて公開"

# オープンデータにて公開の都道府県（令和6年度版の一覧）。収録範囲は dataset.yml・
# README・テーブルの宣言に都道府県名まで書いてあるので、一覧が増えても減っても
# 停止して宣言を直す。判定は完全一致なので、条件の文言が変わった県は減る側に出る
OPEN_PREFECTURES = {
    "01",
    "02",
    "03",
    "07",
    "08",
    "09",
    "10",
    "11",
    "15",
    "20",
    "21",
    "22",
    "25",
    "30",
    "31",
    "32",
    "33",
    "36",
    "39",
    "41",
    "43",
    "46",
    "47",
}

PREFECTURE_CODES = {f"{i:02d}" for i in range(1, 48)}


def _prefectures() -> set[str] | None:
    """処理対象を絞る都道府県コード。

    NLFTP_STEEP_SLOPE_PREFECTURES（カンマ区切り、例: "13,47"）で絞り込める。
    未指定なら公開条件で許された都道府県すべて。
    """
    env = os.environ.get("NLFTP_STEEP_SLOPE_PREFECTURES")
    if not env:
        return None
    return {p.strip() for p in env.split(",") if p.strip()}


def _fetch_page() -> str:
    """配布ページの HTML を取得する（HTML コメントは落とす）。"""
    req = Request(PAGE_URL, headers={"User-Agent": "dataset-nlftp"})
    with urlopen(req) as resp:
        html = resp.read().decode("utf-8", errors="replace")
    return re.sub(r"<!--.*?-->", "", html, flags=re.DOTALL)


def _parse_terms(xlsx_path: Path) -> dict[str, dict[str, str]]:
    """公開条件の一覧を都道府県コード → 各列の辞書にする。

    列の並びが変わったまま読み進めると条件付き公開の県を公開してしまうため、
    見出し行が想定どおりかを先に確かめる。
    """
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    # 1 行目は表題で、2 行目が見出し
    header = rows[1]
    for field, (index, expected) in TERMS_COLUMNS.items():
        actual = str(header[index] or "").replace("\n", "")
        if expected not in actual:
            raise SystemExit(
                f"terms sheet layout changed: column {index} for {field} "
                f"is {actual!r}, expected to contain {expected!r}"
            )

    terms: dict[str, dict[str, str]] = {}
    for row in rows[2:]:
        code = str(row[TERMS_COLUMNS["prefecture_code"][0]] or "").strip()
        if code not in PREFECTURE_CODES:
            continue
        terms[code] = {
            field: str(row[index] or "").strip()
            for field, (index, _) in TERMS_COLUMNS.items()
        }

    missing = sorted(PREFECTURE_CODES - set(terms))
    if missing:
        raise SystemExit(f"prefectures missing from the terms of use: {missing}")
    return terms


def _load_terms(dest: Path) -> dict[str, dict[str, str]]:
    """公開条件の一覧を取得してパースする。"""
    xlsx_path = dest / "terms_of_use.xlsx"
    logger.info("  downloading terms of use...")
    download(TERMS_URL, xlsx_path)
    try:
        terms = _parse_terms(xlsx_path)
    finally:
        xlsx_path.unlink(missing_ok=True)

    open_codes = _open_codes(terms)
    if open_codes != OPEN_PREFECTURES:
        added = sorted(open_codes - OPEN_PREFECTURES)
        removed = sorted(OPEN_PREFECTURES - open_codes)
        raise SystemExit(
            f"terms of use changed: open prefectures added {added}, removed {removed}"
        )
    logger.info(f"  terms of use: {len(open_codes)} prefectures open data")
    return terms


def _open_codes(terms: dict[str, dict[str, str]]) -> set[str]:
    """使用許諾が「オープンデータにて公開」の都道府県コード。"""
    return {c for c, t in terms.items() if t["disclosure"] == OPEN_DISCLOSURE}


def _write_terms(dest: Path, terms: dict[str, dict[str, str]]) -> None:
    """都道府県別の公開条件を Parquet に書き出す。"""
    columns = list(TERMS_COLUMNS)
    con = duckdb.connect()
    try:
        con.execute(
            "CREATE TEMP TABLE terms ("
            + ", ".join(f"{name} VARCHAR" for name in columns)
            + ")"
        )
        con.executemany(
            f"INSERT INTO terms VALUES ({', '.join('?' * len(columns))})",
            [tuple(terms[code][name] for name in columns) for code in sorted(terms)],
        )
        con.execute(
            f"COPY terms TO '{(dest / 'terms.parquet').as_posix()}' "
            "(FORMAT PARQUET, COMPRESSION ZSTD)"
        )
    finally:
        con.close()


def _parse_files(html: str) -> list[tuple[str, str, str]]:
    """ダウンロード一覧から (ファイル名, 都道府県コード, URL) を取り出す。"""
    files = []
    seen = set()
    for path in re.findall(r"DownLd\([^)]*'((?:\.\./)?data/A47/[^']+\.zip)'", html):
        url = urljoin(PAGE_URL, path)
        stem = url.rsplit("/", 1)[-1].removesuffix("_GML.zip")
        match = re.fullmatch(rf"{FILE_PREFIX}_(\d{{2}})", stem)
        if not match or match.group(1) not in PREFECTURE_CODES or stem in seen:
            continue
        seen.add(stem)
        files.append((stem, match.group(1), url))
    # 2021年度整備分があるのは急傾斜地崩壊危険区域を指定している 37 都道府県。
    # 一覧の読み取りが壊れたときに黙って収録が減らないよう下限で止める
    if len(files) < 30:
        raise SystemExit(f"download list parse looks broken: {len(files)} files")
    return sorted(files)


def _extract(zip_path: Path, stem: str, tmp_dir: Path) -> Path:
    """zip から GeoJSON を取り出す。"""
    name = f"{stem}.geojson"
    with zipfile.ZipFile(zip_path) as zf:
        member = next(
            (
                i
                for i in zf.infolist()
                # zip 内のパス区切りが円記号のことがあるので末尾で照合する
                if i.orig_filename.replace("\\", "/").endswith(name)
            ),
            None,
        )
        if member is None:
            raise SystemExit(f"{name} not found in {zip_path.name}")
        path = tmp_dir / name
        with zf.open(member) as src, open(path, "wb") as dst:
            shutil.copyfileobj(src, dst, 1024 * 1024)
    return path


def _convert(geojson: Path, parquet_path: Path, pref: str) -> None:
    """GeoJSON を Parquet に変換する。

    ジオメトリは WKB (BLOB) で保存し、dbt 側で ST_GeomFromWKB で復元する。
    一時ファイルに書き出してからリネームすることで、中断時に不完全な
    Parquet が変換済みとして残らないようにする。
    """
    tmp_path = parquet_path.with_suffix(".parquet.tmp")
    con = duckdb.connect()
    try:
        con.execute("INSTALL spatial; LOAD spatial;")
        con.execute(
            f"""
            COPY (
                SELECT
                    '{pref}' AS prefecture_code,
                    CAST(A47_002 AS VARCHAR) AS municipality_code,
                    CAST(A47_003 AS VARCHAR) AS municipality_name,
                    CAST(A47_004 AS VARCHAR) AS zone_name,
                    CAST(A47_005 AS VARCHAR) AS address,
                    CAST(A47_006 AS VARCHAR) AS notice_date,
                    CAST(A47_007 AS VARCHAR) AS notice_number,
                    CAST(A47_008 AS DOUBLE) AS designated_area_ha,
                    ST_AsWKB(geom) AS geom
                FROM ST_Read('{geojson.as_posix()}')
            ) TO '{tmp_path.as_posix()}' (FORMAT PARQUET, COMPRESSION ZSTD)
            """
        )
    finally:
        con.close()
    tmp_path.rename(parquet_path)


def download_steep_slope(dest_dir: str) -> None:
    """急傾斜地崩壊危険区域データ（オープンデータ公開の都道府県のみ）を Parquet 化する。

    都道府県ごとに zip をダウンロードして Parquet に変換し、zip は変換後すぐ
    削除する。変換済みのファイルはスキップするため、途中で中断しても再実行で
    続きから処理できる（冪等）。
    """
    dest = Path(dest_dir)
    tmp_dir = dest / "tmp"
    parquet_dir = dest / "parquet"
    for d in (tmp_dir, parquet_dir):
        d.mkdir(parents=True, exist_ok=True)

    terms = _load_terms(dest)
    open_codes = _open_codes(terms)
    _write_terms(dest, terms)

    html = _fetch_page()
    files = _parse_files(html)
    logger.info(f"  {len(files)} prefecture files listed")

    only = _prefectures()
    for stem, pref, url in files:
        if pref not in open_codes or (only is not None and pref not in only):
            continue

        parquet_path = parquet_dir / f"{stem}.parquet"
        if parquet_path.exists():
            logger.info(f"  skip {stem} (already converted)")
            continue

        zip_path = tmp_dir / f"{stem}_GML.zip"
        logger.info(f"  downloading {stem}...")
        download(url, zip_path)

        try:
            geojson_path = _extract(zip_path, stem, tmp_dir)
            try:
                _convert(geojson_path, parquet_path, pref)
            finally:
                geojson_path.unlink(missing_ok=True)
        finally:
            zip_path.unlink(missing_ok=True)

    logger.info(f"  steep slope hazard zone data ready in {parquet_dir}")
