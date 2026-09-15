"""国土数値情報 A46 地すべり防止区域データのダウンロード。

国土交通省「国土数値情報ダウンロードサイト」から地すべり防止区域データを
都道府県単位でダウンロードし、Parquet に変換して配置する。地すべり等防止法
第 3 条に基づいて主務大臣が指定した区域で、区域内では地下水を増やす行為や
のり切・切土などが都道府県知事の許可を要する。

利用条件による絞り込み
----------------------
地すべり防止区域は所管省庁が 3 つに分かれており（国土交通省・農林水産省
農村振興局・林野庁）、提供元の配布も都道府県 × 所管省庁の単位になっている。
「各都道府県のデータの利用条件について」（都道府県別の一覧）の使用許諾も
同じ単位で、1 つの都道府県の中で所管省庁ごとに

  オープンデータにて公開 / 条件付き公開 / 非公開 / 未作成

が食い違う。ここでは「オープンデータにて公開」の都道府県 × 所管省庁だけを
取り込む（条件付き公開は成果物への特記事項の明記など県ごとの条件の承諾が
前提になる）。zip には所管省庁ごとの GeoJSON が別ファイルで入っているので、
絞り込みはファイル単位でできる。判定は一覧を毎回取得して行い、列の並びが
変わった場合・一覧に無い都道府県がある場合はエラーで停止する。

配布ファイルの形式
------------------
zip には所管省庁ごとに 1 ファイルの GeoJSON が入る（A46-a が国土交通省、
A46-b が農林水産省農村振興局、A46-c が林野庁）。シェープファイルも同梱される
が、.dbf が CP932 で DuckDB の ST_Read が読めないため GeoJSON を使う。
属性名は所管省庁ごとに接頭辞が変わる（A46-a_001 / A46-b_001 / A46-c_001）。

ダウンロード一覧には 2020 年度整備分と 2021 年度整備分が並ぶ。2021 年度分は
2020 年度分（44 都道府県）を含む 46 都道府県をそろえているので、2021 年度分
だけを取る。一覧には地方単位のまとめ（都道府県コードの範囲外の 52〜59）も
並ぶが、これは都道府県の zip をまとめた zip なので都道府県コードの一致で落ちる。

林野庁所管の区域には点形式で提供された県があるが、オープンデータにて公開の
範囲はすべて面形式だった。面以外が混ざったらモデル側で気付けるよう、
変換では形状を絞り込まない。

データソース: 地すべり防止区域（A46、2021年度）
https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-A46-2021.html
"""

import logging
import os
import re
import shutil
import zipfile
from pathlib import Path
from urllib.parse import urljoin

import duckdb
import openpyxl

from pipelines.download import download, fetch_text

logger = logging.getLogger("pipelines")

PAGE_URL = "https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-A46-2021.html"

# 取り込む整備年度（ファイル名の A46-<yy>）
FILE_PREFIX = "A46-21"

# 都道府県 × 所管省庁別の公開条件（提供元が配布する Excel）
TERMS_URL = "https://nlftp.mlit.go.jp/ksj/gml/codelist/R6_Terms_of_Use_Landslide.xlsx"

# 一覧の列位置（0 始まり）と、その列に入っているはずの見出し
TERMS_COLUMNS = {
    "prefecture_code": (0, "都道府県コード"),
    "prefecture_name": (1, "都道府県"),
    "agency_name": (2, "所管官庁"),
    "disclosure": (3, "使用許諾"),
    "condition": (4, "条件"),
    "condition_detail": (5, "条件詳細"),
}

# 所管省庁と、配布ファイル名に入る記号（A46-a / A46-b / A46-c）の対応
AGENCY_SUFFIX = {
    "国土交通省": "a",
    "農林水産省農村振興局": "b",
    "林野庁": "c",
}

# 取り込む使用許諾。これ以外（条件付き公開・非公開・未作成）は除外する
OPEN_DISCLOSURE = "オープンデータにて公開"

# 使用許諾が「オープンデータにて公開」の都道府県 × 所管省庁（令和6年度版の一覧）。
# 収録範囲は dataset.yml・README・テーブルの宣言に都道府県名まで書いてあるので、
# 一覧が増えても減っても停止して宣言を直す。判定は完全一致なので、条件の文言が
# 変わった組み合わせは減る側に出る
OPEN_PAIRS = {
    ("01", "a"),
    ("02", "a"), ("02", "b"), ("02", "c"),
    ("03", "a"), ("03", "b"), ("03", "c"),
    ("04", "b"),
    ("05", "b"), ("05", "c"),
    ("06", "b"), ("06", "c"),
    ("07", "a"), ("07", "b"), ("07", "c"),
    ("08", "a"), ("08", "c"),
    ("09", "a"), ("09", "c"),
    ("10", "a"), ("10", "b"),
    ("11", "a"),
    ("14", "c"),
    ("15", "a"), ("15", "c"),
    ("16", "b"), ("16", "c"),
    ("17", "b"), ("17", "c"),
    ("19", "b"),
    ("20", "a"),
    ("21", "a"), ("21", "b"),
    ("22", "a"), ("22", "b"), ("22", "c"),
    ("23", "c"),
    ("24", "b"), ("24", "c"),
    ("25", "a"),
    ("26", "b"),
    ("27", "a"), ("27", "c"),
    ("28", "b"), ("28", "c"),
    ("29", "c"),
    ("30", "a"),
    ("31", "a"), ("31", "b"), ("31", "c"),
    ("32", "a"), ("32", "b"), ("32", "c"),
    ("33", "a"),
    ("34", "b"), ("34", "c"),
    ("35", "b"), ("35", "c"),
    ("36", "a"), ("36", "b"), ("36", "c"),
    ("38", "a"), ("38", "b"), ("38", "c"),
    ("39", "a"), ("39", "b"), ("39", "c"),
    ("40", "c"),
    ("41", "a"),
    ("42", "a"),
    ("43", "a"), ("43", "b"), ("43", "c"),
    ("44", "c"),
    ("45", "c"),
    ("46", "a"), ("46", "c"),
    ("47", "a"), ("47", "b"),
}

PREFECTURE_CODES = {f"{i:02d}" for i in range(1, 48)}


def _prefectures() -> set[str] | None:
    """処理対象を絞る都道府県コード。

    NLFTP_LANDSLIDE_PREVENTION_PREFECTURES（カンマ区切り、例: "15,36"）で
    絞り込める。未指定なら公開条件で許された都道府県すべて。
    """
    env = os.environ.get("NLFTP_LANDSLIDE_PREVENTION_PREFECTURES")
    if not env:
        return None
    return {p.strip() for p in env.split(",") if p.strip()}


def _fetch_page() -> str:
    """配布ページの HTML を取得する（HTML コメントは落とす）。"""
    html = fetch_text(PAGE_URL)
    return re.sub(r"<!--.*?-->", "", html, flags=re.DOTALL)


def _parse_terms(xlsx_path: Path) -> list[dict[str, str]]:
    """公開条件の一覧を都道府県 × 所管省庁の行にする。

    都道府県コードと都道府県名は所管省庁 3 行のうち先頭の行にしか入っていない
    ので、直前の値を引き継ぐ。列の並びが変わったまま読み進めると条件付き公開の
    組み合わせを公開してしまうため、見出し行が想定どおりかを先に確かめる。
    """
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    # 1 行目は表題で、2 行目が見出し
    header = rows[1]
    for field, (index, expected) in TERMS_COLUMNS.items():
        actual = str(header[index] or "").replace("\n", "")
        if expected not in actual:
            raise SystemExit(
                f"terms sheet layout changed: column {index} for {field} "
                f"is {actual!r}, expected to contain {expected!r}"
            )

    terms: list[dict[str, str]] = []
    code = name = ""
    for row in rows[2:]:
        if row[TERMS_COLUMNS["prefecture_code"][0]]:
            code = str(row[TERMS_COLUMNS["prefecture_code"][0]]).strip()
            name = str(row[TERMS_COLUMNS["prefecture_name"][0]] or "").strip()
        agency = str(row[TERMS_COLUMNS["agency_name"][0]] or "").strip()
        if code not in PREFECTURE_CODES or agency not in AGENCY_SUFFIX:
            continue
        entry = {
            field: str(row[index] or "").strip()
            for field, (index, _) in TERMS_COLUMNS.items()
        }
        entry["prefecture_code"] = code
        entry["prefecture_name"] = name
        terms.append(entry)

    listed = {t["prefecture_code"] for t in terms}
    missing = sorted(PREFECTURE_CODES - listed)
    if missing:
        raise SystemExit(f"prefectures missing from the terms of use: {missing}")
    return terms


def _load_terms(dest: Path) -> list[dict[str, str]]:
    """公開条件の一覧を取得してパースする。"""
    xlsx_path = dest / "terms_of_use.xlsx"
    logger.info("  downloading terms of use...")
    download(TERMS_URL, xlsx_path)
    try:
        terms = _parse_terms(xlsx_path)
    finally:
        xlsx_path.unlink(missing_ok=True)

    open_pairs = _open_pairs(terms)
    if open_pairs != OPEN_PAIRS:
        added = sorted(open_pairs - OPEN_PAIRS)
        removed = sorted(OPEN_PAIRS - open_pairs)
        raise SystemExit(
            f"terms of use changed: open combinations added {added}, removed {removed}"
        )
    logger.info(
        f"  terms of use: {len(open_pairs)} prefecture x agency combinations open data"
    )
    return terms


def _open_pairs(terms: list[dict[str, str]]) -> set[tuple[str, str]]:
    """使用許諾が「オープンデータにて公開」の（都道府県コード, 所管省庁の記号）。"""
    return {
        (t["prefecture_code"], AGENCY_SUFFIX[t["agency_name"]])
        for t in terms
        if t["disclosure"] == OPEN_DISCLOSURE
    }


def _write_terms(dest: Path, terms: list[dict[str, str]]) -> None:
    """都道府県 × 所管省庁別の公開条件を Parquet に書き出す。"""
    columns = list(TERMS_COLUMNS)
    con = duckdb.connect()
    try:
        con.execute(
            "CREATE TEMP TABLE terms ("
            + ", ".join(f"{name} VARCHAR" for name in columns)
            + ")"
        )
        con.executemany(
            f"INSERT INTO terms VALUES ({', '.join('?' * len(columns))})",
            [tuple(t[name] for name in columns) for t in terms],
        )
        con.execute(
            f"COPY terms TO '{(dest / 'terms.parquet').as_posix()}' "
            "(FORMAT PARQUET, COMPRESSION ZSTD)"
        )
    finally:
        con.close()


def _parse_files(html: str) -> list[tuple[str, str, str]]:
    """ダウンロード一覧から (ファイル名, 都道府県コード, URL) を取り出す。"""
    files = []
    seen = set()
    for path in re.findall(r"DownLd\([^)]*'([^']+\.zip)'", html):
        url = urljoin(PAGE_URL, path)
        stem = url.rsplit("/", 1)[-1].removesuffix("_GML.zip")
        match = re.fullmatch(rf"{FILE_PREFIX}_(\d{{2}})", stem)
        if not match or match.group(1) not in PREFECTURE_CODES or stem in seen:
            continue
        seen.add(stem)
        files.append((stem, match.group(1), url))
    # 2021年度整備分があるのは地すべり防止区域が整備済みの 46 都道府県。
    # 収録範囲は都道府県名まで宣言に書いてあるので、オープンデータ公開の
    # 都道府県の zip が一覧から消えたら止める（下限では、消えたのが公開対象の
    # 県でも件数が足りていれば通ってしまう）
    listed = {pref for _, pref, _ in files}
    missing = sorted({pref for pref, _ in OPEN_PAIRS} - listed)
    if missing:
        raise SystemExit(f"download list is missing open prefectures: {missing}")
    return sorted(files)


def _extract(zip_path: Path, pref: str, agency: str, tmp_dir: Path) -> Path:
    """zip から所管省庁の GeoJSON を取り出す。"""
    name = f"A46-{agency}-21_{pref}.geojson"
    with zipfile.ZipFile(zip_path) as zf:
        member = next(
            (
                i
                for i in zf.infolist()
                # zip 内のパス区切りが円記号のことがあるので末尾で照合する
                if i.orig_filename.replace("\\", "/").endswith(name)
            ),
            None,
        )
        if member is None:
            raise SystemExit(f"{name} not found in {zip_path.name}")
        path = tmp_dir / name
        with zf.open(member) as src, open(path, "wb") as dst:
            shutil.copyfileobj(src, dst, 1024 * 1024)
    return path


def _convert(geojson: Path, parquet_path: Path, pref: str, agency: str) -> None:
    """GeoJSON を Parquet に変換する。

    ジオメトリは WKB (BLOB) で保存し、dbt 側で ST_GeomFromWKB で復元する。
    一時ファイルに書き出してからリネームすることで、中断時に不完全な
    Parquet が変換済みとして残らないようにする。
    """
    prefix = f"A46-{agency}"
    tmp_path = parquet_path.with_suffix(".parquet.tmp")
    con = duckdb.connect()
    try:
        con.execute("INSTALL spatial; LOAD spatial;")
        con.execute(
            f"""
            COPY (
                SELECT
                    '{pref}' AS prefecture_code,
                    CAST("{prefix}_002" AS VARCHAR) AS municipality_code,
                    CAST("{prefix}_003" AS VARCHAR) AS municipality_name,
                    CAST("{prefix}_004" AS VARCHAR) AS zone_name,
                    CAST("{prefix}_005" AS VARCHAR) AS address,
                    CAST("{prefix}_006" AS DATE) AS notice_date,
                    CAST("{prefix}_007" AS VARCHAR) AS notice_number,
                    CAST("{prefix}_008" AS DOUBLE) AS designated_area_ha,
                    CAST("{prefix}_009" AS INTEGER) AS competent_agency_code,
                    ST_AsWKB(geom) AS geom
                FROM ST_Read('{geojson.as_posix()}')
            ) TO '{tmp_path.as_posix()}' (FORMAT PARQUET, COMPRESSION ZSTD)
            """
        )
    finally:
        con.close()
    tmp_path.rename(parquet_path)


def download_landslide_prevention(dest_dir: str) -> None:
    """地すべり防止区域データ（オープンデータ公開の組み合わせのみ）を Parquet 化する。

    都道府県ごとに zip をダウンロードし、その中の所管省庁ごとの GeoJSON を
    Parquet に変換して zip は変換後すぐ削除する。変換済みのファイルはスキップ
    するため、途中で中断しても再実行で続きから処理できる（冪等）。
    """
    dest = Path(dest_dir)
    tmp_dir = dest / "tmp"
    parquet_dir = dest / "parquet"
    for d in (tmp_dir, parquet_dir):
        d.mkdir(parents=True, exist_ok=True)

    terms = _load_terms(dest)
    open_pairs = _open_pairs(terms)
    _write_terms(dest, terms)

    html = _fetch_page()
    files = _parse_files(html)
    logger.info(f"  {len(files)} prefecture files listed")

    only = _prefectures()
    for stem, pref, url in files:
        agencies = sorted(a for p, a in open_pairs if p == pref)
        if not agencies or (only is not None and pref not in only):
            continue

        pending = [
            a for a in agencies if not (parquet_dir / f"{stem}_{a}.parquet").exists()
        ]
        if not pending:
            logger.info(f"  skip {stem} (already converted)")
            continue

        zip_path = tmp_dir / f"{stem}_GML.zip"
        logger.info(f"  downloading {stem}...")
        download(url, zip_path)

        try:
            for agency in pending:
                geojson_path = _extract(zip_path, pref, agency, tmp_dir)
                try:
                    _convert(
                        geojson_path,
                        parquet_dir / f"{stem}_{agency}.parquet",
                        pref,
                        agency,
                    )
                finally:
                    geojson_path.unlink(missing_ok=True)
        finally:
            zip_path.unlink(missing_ok=True)

    logger.info(f"  landslide prevention zone data ready in {parquet_dir}")
