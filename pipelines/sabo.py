"""国土数値情報 A52 砂防指定地データのダウンロード。

国土交通省「国土数値情報ダウンロードサイト」から砂防指定地データを都道府県
単位でダウンロードし、Parquet に変換して配置する。砂防法第 2 条に基づいて
国土交通大臣が指定した土地で、指定地内では土石の採取・工作物の設置などが
都道府県の条例で制限される。

配布の範囲
----------
使用許諾は CC BY 4.0 で、都道府県ごとの公開条件による絞り込みは要らない。
代わりに配布があるのが原典資料の提供を受けた 24 都道府県だけで、残りの県は
ダウンロード一覧に並ばない。都道府県ごとのデータ時点は提供元が別の Excel
（各都道府県別のデータ時点）で配布しているので、これも取り込んで
テーブルに載せる。

配布ファイルの形式
------------------
zip には都道府県 1 ファイルの GeoJSON が入る。提供元の仕様では面と線の
2 形状があるが、2023 年度版で配布されているのは面（Polygon）だけだった。
シェープファイルも同梱されるが、.dbf が CP932 で DuckDB の ST_Read が
読めないため GeoJSON を使う。

データソース: 砂防指定地（A52、2023年度）
https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-A52-2023.html
"""

import datetime
import logging
import os
import re
import shutil
import zipfile
from pathlib import Path
from urllib.parse import urljoin

import duckdb
import openpyxl

from pipelines.download import download, fetch_text

logger = logging.getLogger("pipelines")

PAGE_URL = "https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-A52-2023.html"

# 取り込む整備年度（ファイル名の A52-<yy>）
FILE_PREFIX = "A52-23"

# 都道府県別のデータ時点（提供元が配布する Excel）
DATAPOINT_URL = "https://nlftp.mlit.go.jp/ksj/gml/codelist/A52_R5_datapoint.xlsx"

# 一覧の列位置（0 始まり）と、その列に入っているはずの見出し
DATAPOINT_COLUMNS = {
    "prefecture_name": (1, "都道府県名"),
    "as_of": (2, "時点"),
}

# データ時点が無い都道府県を表す記号
DATAPOINT_NONE = "－"

# 配布がある都道府県（2023年度版）。収録範囲は dataset.yml・README・テーブルの
# 宣言に都道府県名まで書いてあるので、一覧が増えても減っても停止して宣言を直す
DISTRIBUTED_PREFECTURES = {
    "02", "03", "08", "09", "11", "12", "13", "14",
    "15", "17", "20", "21", "22", "25", "30", "32",
    "34", "36", "37", "39", "42", "43", "46", "47",
}

PREFECTURE_CODES = {f"{i:02d}" for i in range(1, 48)}

# 都道府県名 → 都道府県コード（データ時点の一覧は名前しか持たない）
PREFECTURE_NAMES = {
    "北海道": "01", "青森県": "02", "岩手県": "03", "宮城県": "04",
    "秋田県": "05", "山形県": "06", "福島県": "07", "茨城県": "08",
    "栃木県": "09", "群馬県": "10", "埼玉県": "11", "千葉県": "12",
    "東京都": "13", "神奈川県": "14", "新潟県": "15", "富山県": "16",
    "石川県": "17", "福井県": "18", "山梨県": "19", "長野県": "20",
    "岐阜県": "21", "静岡県": "22", "愛知県": "23", "三重県": "24",
    "滋賀県": "25", "京都府": "26", "大阪府": "27", "兵庫県": "28",
    "奈良県": "29", "和歌山県": "30", "鳥取県": "31", "島根県": "32",
    "岡山県": "33", "広島県": "34", "山口県": "35", "徳島県": "36",
    "香川県": "37", "愛媛県": "38", "高知県": "39", "福岡県": "40",
    "佐賀県": "41", "長崎県": "42", "熊本県": "43", "大分県": "44",
    "宮崎県": "45", "鹿児島県": "46", "沖縄県": "47",
}


def _prefectures() -> set[str] | None:
    """処理対象を絞る都道府県コード。

    NLFTP_SABO_PREFECTURES（カンマ区切り、例: "15,20"）で絞り込める。
    未指定なら配布のあるすべての都道府県。
    """
    env = os.environ.get("NLFTP_SABO_PREFECTURES")
    if not env:
        return None
    return {p.strip() for p in env.split(",") if p.strip()}


def _fetch_page() -> str:
    """配布ページの HTML を取得する（HTML コメントは落とす）。"""
    html = fetch_text(PAGE_URL)
    return re.sub(r"<!--.*?-->", "", html, flags=re.DOTALL)


def _cell_text(value: object) -> str:
    """データ時点のセルを文字列にする。

    表記は「R5.3末」のような提供元の略記だが、2 県だけ日付セルとして
    入っているので、その 2 件は Excel の日付シリアル値を ISO の日付に直す。
    """
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        epoch = datetime.date(1899, 12, 30)
        return (epoch + datetime.timedelta(days=int(value))).isoformat()
    return str(value).strip()


def _parse_datapoint(xlsx_path: Path) -> dict[str, str]:
    """データ時点の一覧を都道府県コード → 時点の表記にする。"""
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    # 表の前に注意事項が並ぶので、見出し行を探してから読む
    header_index = next(
        (
            i
            for i, row in enumerate(rows)
            if all(
                expected in str(row[index] or "")
                for index, expected in DATAPOINT_COLUMNS.values()
            )
        ),
        None,
    )
    if header_index is None:
        raise SystemExit("data point sheet layout changed: header row not found")

    datapoint: dict[str, str] = {}
    for row in rows[header_index + 1 :]:
        name = str(row[DATAPOINT_COLUMNS["prefecture_name"][0]] or "").strip()
        code = PREFECTURE_NAMES.get(name)
        if code is None:
            continue
        as_of = _cell_text(row[DATAPOINT_COLUMNS["as_of"][0]])
        datapoint[code] = "" if as_of == DATAPOINT_NONE else as_of

    missing = sorted(PREFECTURE_CODES - set(datapoint))
    if missing:
        raise SystemExit(f"prefectures missing from the data point list: {missing}")
    return datapoint


def _load_datapoint(dest: Path) -> dict[str, str]:
    """データ時点の一覧を取得してパースする。"""
    xlsx_path = dest / "datapoint.xlsx"
    logger.info("  downloading data point list...")
    download(DATAPOINT_URL, xlsx_path)
    try:
        datapoint = _parse_datapoint(xlsx_path)
    finally:
        xlsx_path.unlink(missing_ok=True)

    dated = {code for code, as_of in datapoint.items() if as_of}
    if dated != DISTRIBUTED_PREFECTURES:
        added = sorted(dated - DISTRIBUTED_PREFECTURES)
        removed = sorted(DISTRIBUTED_PREFECTURES - dated)
        raise SystemExit(
            f"data point list changed: prefectures added {added}, removed {removed}"
        )
    logger.info(f"  data point list: {len(dated)} prefectures")
    return datapoint


def _write_datapoint(dest: Path, datapoint: dict[str, str]) -> None:
    """都道府県別のデータ時点を Parquet に書き出す。"""
    con = duckdb.connect()
    try:
        con.execute(
            "CREATE TEMP TABLE datapoint (prefecture_code VARCHAR, as_of VARCHAR)"
        )
        con.executemany(
            "INSERT INTO datapoint VALUES (?, ?)",
            [(code, datapoint[code]) for code in sorted(datapoint)],
        )
        con.execute(
            f"COPY datapoint TO '{(dest / 'datapoint.parquet').as_posix()}' "
            "(FORMAT PARQUET, COMPRESSION ZSTD)"
        )
    finally:
        con.close()


def _parse_files(html: str) -> list[tuple[str, str, str]]:
    """ダウンロード一覧から (ファイル名, 都道府県コード, URL) を取り出す。"""
    files = []
    seen = set()
    for path in re.findall(r"DownLd\([^)]*'([^']+\.zip)'", html):
        url = urljoin(PAGE_URL, path)
        stem = url.rsplit("/", 1)[-1].removesuffix("_GML.zip")
        match = re.fullmatch(rf"{FILE_PREFIX}_(\d{{2}})", stem)
        if not match or match.group(1) not in PREFECTURE_CODES or stem in seen:
            continue
        seen.add(stem)
        files.append((stem, match.group(1), url))
    listed = {pref for _, pref, _ in files}
    if listed != DISTRIBUTED_PREFECTURES:
        added = sorted(listed - DISTRIBUTED_PREFECTURES)
        removed = sorted(DISTRIBUTED_PREFECTURES - listed)
        raise SystemExit(
            f"download list changed: prefectures added {added}, removed {removed}"
        )
    return sorted(files)


def _extract(zip_path: Path, stem: str, tmp_dir: Path) -> Path:
    """zip から面形式の GeoJSON を取り出す。"""
    name = f"{stem}Polygon.geojson"
    with zipfile.ZipFile(zip_path) as zf:
        member = next(
            (
                i
                for i in zf.infolist()
                # zip 内のパス区切りが円記号のことがあるので末尾で照合する
                if i.orig_filename.replace("\\", "/").endswith(name)
            ),
            None,
        )
        if member is None:
            raise SystemExit(f"{name} not found in {zip_path.name}")
        path = tmp_dir / name
        with zf.open(member) as src, open(path, "wb") as dst:
            shutil.copyfileobj(src, dst, 1024 * 1024)
    return path


def _convert(geojson: Path, parquet_path: Path, pref: str) -> None:
    """GeoJSON を Parquet に変換する。

    ジオメトリは WKB (BLOB) で保存し、dbt 側で ST_GeomFromWKB で復元する。
    一時ファイルに書き出してからリネームすることで、中断時に不完全な
    Parquet が変換済みとして残らないようにする。
    """
    tmp_path = parquet_path.with_suffix(".parquet.tmp")
    con = duckdb.connect()
    try:
        con.execute("INSTALL spatial; LOAD spatial;")
        con.execute(
            f"""
            COPY (
                SELECT
                    '{pref}' AS prefecture_code,
                    CAST(A52_002 AS VARCHAR) AS municipality_code,
                    CAST(A52_003 AS VARCHAR) AS municipality_name,
                    CAST(A52_004 AS VARCHAR) AS river_name,
                    CAST(A52_005 AS VARCHAR) AS tributary_name,
                    CAST(A52_006 AS VARCHAR) AS notice_date_text,
                    CAST(A52_007 AS VARCHAR) AS notice_number,
                    CAST(A52_008 AS DOUBLE) AS designated_area_ha,
                    CAST(A52_009 AS VARCHAR) AS reference_number,
                    CAST(A52_010 AS VARCHAR) AS designation_method,
                    ST_AsWKB(geom) AS geom
                FROM ST_Read('{geojson.as_posix()}')
            ) TO '{tmp_path.as_posix()}' (FORMAT PARQUET, COMPRESSION ZSTD)
            """
        )
    finally:
        con.close()
    tmp_path.rename(parquet_path)


def download_sabo(dest_dir: str) -> None:
    """砂防指定地データを Parquet 化する。

    都道府県ごとに zip をダウンロードして Parquet に変換し、zip は変換後すぐ
    削除する。変換済みのファイルはスキップするため、途中で中断しても再実行で
    続きから処理できる（冪等）。
    """
    dest = Path(dest_dir)
    tmp_dir = dest / "tmp"
    parquet_dir = dest / "parquet"
    for d in (tmp_dir, parquet_dir):
        d.mkdir(parents=True, exist_ok=True)

    datapoint = _load_datapoint(dest)
    _write_datapoint(dest, datapoint)

    html = _fetch_page()
    files = _parse_files(html)
    logger.info(f"  {len(files)} prefecture files listed")

    only = _prefectures()
    for stem, pref, url in files:
        if only is not None and pref not in only:
            continue

        parquet_path = parquet_dir / f"{stem}.parquet"
        if parquet_path.exists():
            logger.info(f"  skip {stem} (already converted)")
            continue

        zip_path = tmp_dir / f"{stem}_GML.zip"
        logger.info(f"  downloading {stem}...")
        download(url, zip_path)

        try:
            geojson_path = _extract(zip_path, stem, tmp_dir)
            try:
                _convert(geojson_path, parquet_path, pref)
            finally:
                geojson_path.unlink(missing_ok=True)
        finally:
            zip_path.unlink(missing_ok=True)

    logger.info(f"  sabo designated area data ready in {parquet_dir}")
