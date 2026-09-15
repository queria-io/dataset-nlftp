"""国土数値情報 A48 災害危険区域データのダウンロード。

国土交通省「国土数値情報ダウンロードサイト」から災害危険区域データを都道府県
単位でダウンロードし、Parquet に変換して配置する。建築基準法第 39 条に基づいて
地方公共団体が条例で指定する区域で、区域内では住居の用に供する建築物の建築が
条例の定めるところにより制限される。

指定するのは都道府県と市町村の両方で、砂防三法（A46 地すべり防止区域・
A47 急傾斜地崩壊危険区域・A52 砂防指定地）のように都道府県知事が指定する区域とは
指定の主体が違う。

利用条件による絞り込み
----------------------
災害危険区域データは原典資料を作成した地方公共団体ごとに公開条件が異なる。
提供元が配布する「各地方公共団体のデータ使用条件について」には使用許諾が

  オープンデータにて公開 / 条件付き公開 / 非公開 / 未作成

のいずれかで示され、1 行が 1 つの指定公共団体（県または市町村）にあたる。
ここでは「オープンデータにて公開」の指定公共団体だけを取り込む。条件付き公開の
条件詳細は「商用利用不可・再配布不可」から精度の注意書きまで団体ごとにばらばらで、
機械的に読み分けられないため区分ごと落とす。

一覧の行と地物の対応は指定主体区分（A48_004）で決まる。区分が都道府県なら
その都道府県の県の行（行政CD が <都道府県コード>000）、市町村なら代表行政
コード（A48_003）の行を見る。判定は一覧を毎回取得して行い、列の並びが変わった
場合・オープンデータにて公開の集合が変わった場合はエラーで停止する。

指定主体区分は原典の値をそのまま使うが、市が指定した区域に都道府県が入っている
ことがある（愛媛県は全行が都道府県で、うち 283 行は根拠条例が松山市・今治市・
新居浜市・西条市の条例）。今の配布ではその 4 市とも一覧でオープンデータにて公開
なので絞り込みの結果は変わらないが、条件付き公開・非公開の団体の条例で指定された
区域が同じ形で紛れ込むと区分だけでは気づけない。取り込んだ区域の根拠条例に、
同じ都道府県の一覧でオープンデータにて公開ではない市町村の名前が現れたら停止する。

一覧に無い代表行政コードが 2 件ある（長崎県の時津町・佐々町）。いずれも根拠条例が
長崎市・佐世保市の条例で、区域の所在地と指定した市町村が食い違っている。どちらの
行を見るべきか機械的に決められないので取り込まない。想定外のコードが増えたら
停止する。

配布ファイルの形式
------------------
zip には都道府県 1 ファイルの GeoJSON が入る。シェープファイルも同梱されるが、
.dbf が CP932 で DuckDB の ST_Read が読めないため GeoJSON を使う。

形状は面と点の 2 つがあるが、点（A48P-21_34.geojson）が配布されるのは広島県だけで、
広島県の県指定の区域は条件付き公開にあたる。取り込める点の地物が無いので面だけを
扱う。山形県は区域の属性だけがあり形状のデータが配布されていない。

ダウンロード一覧には 2020 年度整備分と 2021 年度整備分が並ぶ。2021 年度分は同じ
都道府県の作り直しで、2020 年度分のポリゴン 14,391 のうち 2021 年度分のどの
ポリゴンとも交わらないものは 5（岩手 1・富山 3・愛媛 1）しかない。2021 年度分
だけを取り、2020 年度分は取らない。一覧には地方単位のまとめ（都道府県コードの
範囲外の 52〜59）も並ぶが、これは都道府県の zip をまとめた zip なので都道府県
コードの一致で落ちる。

データソース: 災害危険区域（A48、2021年度）
https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-A48-2021.html
"""

import logging
import os
import re
import shutil
import zipfile
from pathlib import Path
from urllib.parse import urljoin
from urllib.request import Request, urlopen

import duckdb
import openpyxl

from pipelines.download import download

logger = logging.getLogger("pipelines")

PAGE_URL = "https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-A48-2021.html"

# 取り込む整備年度（ファイル名の A48-<yy>）
FILE_PREFIX = "A48-21"

# 地方公共団体別の公開条件（提供元が配布する Excel）
TERMS_URL = (
    "https://nlftp.mlit.go.jp/ksj/gml/codelist/R6_Terms_of_Use_DisasterRiskArea.xlsx"
)

# 一覧の列位置（0 始まり）と、その列に入っているはずの見出し
TERMS_COLUMNS = {
    "prefecture_code": (0, "都道府県"),
    "prefecture_name": (1, "都道府県"),
    "organization": (2, "指定公共団体"),
    "reason": (3, "指定理由"),
    "admin_code": (4, "行政CD"),
    "disclosure": (5, "使用許諾"),
    "condition": (6, "条件"),
    "condition_detail": (7, "条件詳細"),
}

# 取り込む使用許諾。これ以外（条件付き公開・非公開・未作成）は除外する
OPEN_DISCLOSURE = "オープンデータにて公開"

# オープンデータにて公開の行政CD（令和6年度版の一覧）。末尾が 000 のものは
# 都道府県が指定した区域、それ以外は市町村が指定した区域にあたる。収録範囲は
# dataset.yml・README・テーブルの宣言に書いてあるので、一覧が増えても減っても
# 停止して宣言を直す。判定は完全一致なので、条件の文言が変わった団体は減る側に出る
OPEN_ADMIN_CODES = {
    "01100",
    "01202",
    "01367",
    "01584",
    "01663",
    "02000",
    "02202",
    "02203",
    "02204",
    "03202",
    "03203",
    "03210",
    "03211",
    "03461",
    "03482",
    "03483",
    "03503",
    "04100",
    "04202",
    "04203",
    "04205",
    "04207",
    "04211",
    "04214",
    "04341",
    "04361",
    "04362",
    "04404",
    "04581",
    "04606",
    "05201",
    "05212",
    "06000",
    "07204",
    "07208",
    "07209",
    "07210",
    "07542",
    "07547",
    "07561",
    "08000",
    "08201",
    "08202",
    "08203",
    "08212",
    "08214",
    "08215",
    "08217",
    "08220",
    "08221",
    "09000",
    "10000",
    "12000",
    "13381",
    "14000",
    "14100",
    "14130",
    "14150",
    "14201",
    "14203",
    "14204",
    "14205",
    "14206",
    "14211",
    "14212",
    "15000",
    "16000",
    "17000",
    "18000",
    "18204",
    "20211",
    "21000",
    "23100",
    "25000",
    "26201",
    "26202",
    "27000",
    "27220",
    "28000",
    "28219",
    "30000",
    "30203",
    "30206",
    "30207",
    "31000",
    "32000",
    "32207",
    "32441",
    "32448",
    "32449",
    "34214",
    "36000",
    "36206",
    "36468",
    "38000",
    "38201",
    "38202",
    "38205",
    "38206",
    "41000",
    "42000",
    "42201",
    "42202",
    "43000",
    "43214",
    "43215",
    "43348",
    "43425",
    "43444",
    "45201",
    "45206",
    "46000",
    "46215",
    "46392",
    "47000",
}

# 配布データにあるが一覧に無い代表行政コード。判定できないので取り込まない
UNLISTED_ADMIN_CODES = {"42308", "42391"}

# 形状のデータが配布されていない都道府県（属性だけの一覧が入っている）
NO_GEOMETRY_PREFECTURES = {"06"}

# 点の GeoJSON が配布されている都道府県。県指定の区域が条件付き公開にあたるため
# 取り込める点の地物が無く、面だけを扱う
POINT_PREFECTURES = {"34"}

PREFECTURE_CODES = {f"{i:02d}" for i in range(1, 48)}


def _prefectures() -> set[str] | None:
    """処理対象を絞る都道府県コード。

    NLFTP_DISASTER_RISK_AREA_PREFECTURES（カンマ区切り、例: "13,47"）で絞り込める。
    未指定なら配布されている都道府県すべて。
    """
    env = os.environ.get("NLFTP_DISASTER_RISK_AREA_PREFECTURES")
    if not env:
        return None
    return {p.strip() for p in env.split(",") if p.strip()}


def _fetch_page() -> str:
    """配布ページの HTML を取得する（HTML コメントは落とす）。"""
    req = Request(PAGE_URL, headers={"User-Agent": "dataset-nlftp"})
    with urlopen(req) as resp:
        html = resp.read().decode("utf-8", errors="replace")
    return re.sub(r"<!--.*?-->", "", html, flags=re.DOTALL)


def _parse_terms(xlsx_path: Path) -> list[dict[str, str]]:
    """公開条件の一覧を行ごとの辞書にする。

    列の並びが変わったまま読み進めると条件付き公開の団体を公開してしまうため、
    見出し行が想定どおりかを先に確かめる。都道府県コードと都道府県名はセルが
    縦に結合されていて 2 行目以降が空になるので、直前の値で埋める。
    """
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    # 1 行目は表題で、2 行目が見出し
    header = rows[1]
    for field, (index, expected) in TERMS_COLUMNS.items():
        actual = str(header[index] or "").replace("\n", "")
        if expected not in actual:
            raise SystemExit(
                f"terms sheet layout changed: column {index} for {field} "
                f"is {actual!r}, expected to contain {expected!r}"
            )

    terms: list[dict[str, str]] = []
    prefecture_code = ""
    prefecture_name = ""
    for row in rows[2:]:
        if all(value is None for value in row):
            continue
        if row[TERMS_COLUMNS["prefecture_code"][0]]:
            prefecture_code = str(row[TERMS_COLUMNS["prefecture_code"][0]]).strip()
            prefecture_name = str(
                row[TERMS_COLUMNS["prefecture_name"][0]] or ""
            ).strip()
        entry = {
            field: str(row[index] or "").strip()
            for field, (index, _) in TERMS_COLUMNS.items()
        }
        entry["prefecture_code"] = prefecture_code
        entry["prefecture_name"] = prefecture_name
        terms.append(entry)

    listed = {t["prefecture_code"] for t in terms}
    missing = sorted(PREFECTURE_CODES - listed)
    if missing:
        raise SystemExit(f"prefectures missing from the terms of use: {missing}")
    return terms


def _open_codes(terms: list[dict[str, str]]) -> set[str]:
    """使用許諾が「オープンデータにて公開」の行政CD。"""
    return {
        t["admin_code"]
        for t in terms
        if t["admin_code"] and t["disclosure"] == OPEN_DISCLOSURE
    }


def _restricted_names(terms: list[dict[str, str]]) -> dict[str, list[str]]:
    """都道府県コード → オープンデータにて公開ではない市町村の名前。

    根拠条例の照合に使うので、指定公共団体が市町村として書かれている行だけを
    拾う（「県」「府」のような都道府県の行と、直前の行の繰り返しを表す「〃」は
    条例名の一部として現れないか、現れても総当たりになる）。
    """
    names: dict[str, list[str]] = {}
    for t in terms:
        if t["disclosure"] == OPEN_DISCLOSURE:
            continue
        name = t["organization"].replace("〃", "").strip()
        if not name.endswith(("市", "区", "町", "村")):
            continue
        names.setdefault(t["prefecture_code"], []).append(name)
    return names


def _load_terms(dest: Path) -> list[dict[str, str]]:
    """公開条件の一覧を取得してパースする。"""
    xlsx_path = dest / "terms_of_use.xlsx"
    logger.info("  downloading terms of use...")
    download(TERMS_URL, xlsx_path)
    try:
        terms = _parse_terms(xlsx_path)
    finally:
        xlsx_path.unlink(missing_ok=True)

    open_codes = _open_codes(terms)
    if open_codes != OPEN_ADMIN_CODES:
        added = sorted(open_codes - OPEN_ADMIN_CODES)
        removed = sorted(OPEN_ADMIN_CODES - open_codes)
        raise SystemExit(
            f"terms of use changed: open organizations added {added}, removed {removed}"
        )
    logger.info(f"  terms of use: {len(open_codes)} organizations open data")
    return terms


def _write_terms(dest: Path, terms: list[dict[str, str]]) -> None:
    """地方公共団体別の公開条件を Parquet に書き出す。"""
    columns = list(TERMS_COLUMNS)
    con = duckdb.connect()
    try:
        con.execute(
            "CREATE TEMP TABLE terms ("
            + ", ".join(f"{name} VARCHAR" for name in columns)
            + ")"
        )
        con.executemany(
            f"INSERT INTO terms VALUES ({', '.join('?' * len(columns))})",
            [tuple(t[name] for name in columns) for t in terms],
        )
        con.execute(
            f"COPY terms TO '{(dest / 'terms.parquet').as_posix()}' "
            "(FORMAT PARQUET, COMPRESSION ZSTD)"
        )
    finally:
        con.close()


def _parse_files(html: str) -> list[tuple[str, str, str]]:
    """ダウンロード一覧から (ファイル名, 都道府県コード, URL) を取り出す。"""
    files = []
    seen = set()
    for path in re.findall(r"DownLd\([^)]*'((?:\.\./)?data/A48/[^']+\.zip)'", html):
        url = urljoin(PAGE_URL, path)
        stem = url.rsplit("/", 1)[-1].removesuffix("_GML.zip")
        match = re.fullmatch(rf"{FILE_PREFIX}_(\d{{2}})", stem)
        if not match or match.group(1) not in PREFECTURE_CODES or stem in seen:
            continue
        seen.add(stem)
        files.append((stem, match.group(1), url))
    # 2021年度整備分があるのは災害危険区域を指定している 37 都道府県。
    # 一覧の読み取りが壊れたときに黙って収録が減らないよう下限で止める
    if len(files) < 30:
        raise SystemExit(f"download list parse looks broken: {len(files)} files")
    return sorted(files)


def _extract(zip_path: Path, stem: str, pref: str, tmp_dir: Path) -> Path | None:
    """zip から面の GeoJSON を取り出す。

    形状のデータが無い都道府県では None を返す。点の GeoJSON は取り込まないが、
    配布される都道府県が変わったら止める。
    """
    name = f"{stem}.geojson"
    with zipfile.ZipFile(zip_path) as zf:
        members = {
            # zip 内のパス区切りが円記号のことがあるので末尾で照合する
            i.orig_filename.replace("\\", "/").rsplit("/", 1)[-1]: i
            for i in zf.infolist()
            if i.orig_filename.lower().endswith(".geojson")
        }
        others = set(members) - {name}
        if others and pref not in POINT_PREFECTURES:
            raise SystemExit(f"unexpected geojson in {zip_path.name}: {sorted(others)}")
        member = members.get(name)
        if member is None:
            if pref in NO_GEOMETRY_PREFECTURES:
                logger.info(f"  skip {stem} (no geometry distributed)")
                return None
            raise SystemExit(f"{name} not found in {zip_path.name}")
        path = tmp_dir / name
        with zf.open(member) as src, open(path, "wb") as dst:
            shutil.copyfileobj(src, dst, 1024 * 1024)
    return path


def _convert(
    geojson: Path,
    parquet_path: Path,
    pref: str,
    listed: set[str],
    restricted: list[str],
) -> None:
    """GeoJSON を公開条件で絞り込んで Parquet に変換する。

    ジオメトリは WKB (BLOB) で保存し、dbt 側で ST_GeomFromWKB で復元する。
    一時ファイルに書き出してからリネームすることで、中断時に不完全な
    Parquet が変換済みとして残らないようにする。
    """
    tmp_path = parquet_path.with_suffix(".parquet.tmp")
    allowed = ", ".join(f"'{code}'" for code in sorted(OPEN_ADMIN_CODES))
    known = ", ".join(f"'{code}'" for code in sorted(listed | UNLISTED_ADMIN_CODES))
    con = duckdb.connect()
    try:
        con.execute("INSTALL spatial; LOAD spatial;")
        con.execute(
            f"""
            CREATE TEMP TABLE source AS
            SELECT
                '{pref}' AS prefecture_code,
                CAST(A48_003 AS VARCHAR) AS municipality_code,
                CAST(A48_002 AS VARCHAR) AS municipality_name,
                CAST(A48_004 AS VARCHAR) AS designating_body_code,
                CAST(A48_005 AS VARCHAR) AS zone_name,
                CAST(A48_006 AS VARCHAR) AS address,
                CAST(A48_007 AS VARCHAR) AS reason_code,
                CAST(A48_008 AS VARCHAR) AS reason_detail,
                CAST(A48_009 AS VARCHAR) AS notice_date,
                CAST(A48_010 AS VARCHAR) AS notice_number,
                CAST(A48_011 AS VARCHAR) AS ordinance_name,
                CAST(A48_012 AS DOUBLE) AS designated_area_ha,
                CAST(A48_013 AS VARCHAR) AS reference_scale,
                CAST(A48_014 AS VARCHAR) AS remarks,
                ST_AsWKB(geom) AS geom,
                CASE
                    WHEN CAST(A48_004 AS VARCHAR) = '1' THEN '{pref}000'
                    ELSE CAST(A48_003 AS VARCHAR)
                END AS admin_code
            FROM ST_Read('{geojson.as_posix()}')
            """
        )
        # 一覧に無い行政コードが増えたら、公開条件を引けず取り込みの可否を
        # 決められないので止める
        unknown = [
            row[0]
            for row in con.execute(
                f"SELECT DISTINCT admin_code FROM source "
                f"WHERE admin_code NOT IN ({known}) ORDER BY 1"
            ).fetchall()
        ]
        if unknown:
            raise SystemExit(
                f"{geojson.name}: admin codes not in the terms of use: {unknown}"
            )
        # 指定主体区分が原典で都道府県になっている市町村の区域があるため、区分
        # だけでは条件付き公開・非公開の団体の区域を見分けられない。取り込む行の
        # 根拠条例に、同じ都道府県でオープンデータにて公開ではない市町村の名前が
        # 現れたら止める
        for name in restricted:
            found = con.execute(
                f"SELECT count(*) FROM source "
                f"WHERE admin_code IN ({allowed}) AND ordinance_name LIKE ?",
                [f"%{name}%"],
            ).fetchone()[0]
            if found:
                raise SystemExit(
                    f"{geojson.name}: {found} features are designated under "
                    f"{name}'s ordinance, which is not open data"
                )

        total, kept = con.execute(
            f"SELECT count(*), count(*) FILTER (WHERE admin_code IN ({allowed})) "
            "FROM source"
        ).fetchone()
        con.execute(
            f"""
            COPY (
                SELECT * EXCLUDE (admin_code) FROM source
                WHERE admin_code IN ({allowed})
            ) TO '{tmp_path.as_posix()}' (FORMAT PARQUET, COMPRESSION ZSTD)
            """
        )
    finally:
        con.close()
    logger.info(f"  {parquet_path.stem}: {kept}/{total} features open data")
    tmp_path.rename(parquet_path)


def download_disaster_risk_area(dest_dir: str) -> None:
    """災害危険区域データ（オープンデータ公開の地方公共団体のみ）を Parquet 化する。

    都道府県ごとに zip をダウンロードして Parquet に変換し、zip は変換後すぐ
    削除する。変換済みのファイルはスキップするため、途中で中断しても再実行で
    続きから処理できる（冪等）。
    """
    dest = Path(dest_dir)
    tmp_dir = dest / "tmp"
    parquet_dir = dest / "parquet"
    for d in (tmp_dir, parquet_dir):
        d.mkdir(parents=True, exist_ok=True)

    terms = _load_terms(dest)
    _write_terms(dest, terms)
    listed = {t["admin_code"] for t in terms if t["admin_code"]}
    restricted = _restricted_names(terms)

    html = _fetch_page()
    files = _parse_files(html)
    logger.info(f"  {len(files)} prefecture files listed")

    only = _prefectures()
    for stem, pref, url in files:
        if only is not None and pref not in only:
            continue

        parquet_path = parquet_dir / f"{stem}.parquet"
        if parquet_path.exists():
            logger.info(f"  skip {stem} (already converted)")
            continue

        zip_path = tmp_dir / f"{stem}_GML.zip"
        logger.info(f"  downloading {stem}...")
        download(url, zip_path)

        try:
            geojson_path = _extract(zip_path, stem, pref, tmp_dir)
            if geojson_path is None:
                continue
            try:
                _convert(
                    geojson_path,
                    parquet_path,
                    pref,
                    listed,
                    restricted.get(pref, []),
                )
            finally:
                geojson_path.unlink(missing_ok=True)
        finally:
            zip_path.unlink(missing_ok=True)

    logger.info(f"  disaster risk area data ready in {parquet_dir}")
